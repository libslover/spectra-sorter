"""
GUI интерфейс на tkinter для работы с Excel файлами и протоколами спектров.
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
from typing import Optional, List, Dict, Any
import json
import os

from excel_handler import ExcelWorkbook, ExcelWorksheet, RowValidationError
from openpyxl.styles import Font, PatternFill, Alignment


class ProtocolManager:
    """Управление протоколами (JSON файлы)."""
    
    def __init__(self, protocol_dir: str = "."):
        self.protocol_dir = protocol_dir
        self.current_protocol: Optional[Dict] = None
        self.current_protocol_path: Optional[str] = None
    
    def get_new_protocol_name(self) -> str:
        """Получить новое имя протокола с автоматической нумерацией."""
        base_name = "Протокол1"
        if not os.path.exists(os.path.join(self.protocol_dir, f"{base_name}.json")):
            return base_name
        
        counter = 2
        while True:
            name = f"Протокол{counter}"
            if not os.path.exists(os.path.join(self.protocol_dir, f"{name}.json")):
                return name
            counter += 1
    
    def create_protocol(self, name: Optional[str] = None) -> Dict:
        """Создать новый протокол."""
        if name is None:
            name = self.get_new_protocol_name()
        
        protocol = {
            "name": name,
            "ranges": [],
            "lines": []
        }
        
        path = os.path.join(self.protocol_dir, f"{name}.json")
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(protocol, f, ensure_ascii=False, indent=2)
        
        self.current_protocol = protocol
        self.current_protocol_path = path
        return protocol
    
    def open_protocol(self, path: str) -> Dict:
        """Открыть существующий протокол."""
        with open(path, 'r', encoding='utf-8') as f:
            protocol = json.load(f)
        
        self.current_protocol = protocol
        self.current_protocol_path = path
        return protocol
    
    def save_protocol(self):
        """Сохранить текущий протокол."""
        if not self.current_protocol or not self.current_protocol_path:
            raise ValueError("Нет открытого протокола")
        
        with open(self.current_protocol_path, 'w', encoding='utf-8') as f:
            json.dump(self.current_protocol, f, ensure_ascii=False, indent=2)
    
    def add_range(self, name: str):
        """Добавить диапазон в протокол."""
        if not self.current_protocol:
            raise ValueError("Нет открытого протокола")
        
        self.current_protocol["ranges"].append({"name": name})
        self.save_protocol()
    
    def add_line(self, range_name: str, particle: str, w0: float):
        """Добавить линию в протокол."""
        if not self.current_protocol:
            raise ValueError("Нет открытого протокола")
        
        line = {
            "range": range_name,
            "particle": particle,
            "w0": w0
        }
        self.current_protocol["lines"].append(line)
        self.save_protocol()
    
    def find_nearest_line(self, w: float, max_distance: float = 1.0) -> Optional[Dict]:
        """Найти ближайшую линию по значению w."""
        if not self.current_protocol or not self.current_protocol["lines"]:
            return None
        
        nearest = None
        min_dist = float('inf')
        
        for line in self.current_protocol["lines"]:
            dist = abs(line["w0"] - w)
            if dist < min_dist:
                min_dist = dist
                nearest = line
        
        if nearest and min_dist <= max_distance:
            return nearest
        
        return None


class SpectrumExcelApp:
    """Приложение для работы с файлами спектров в Excel."""
    
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("Обработка спектров - Excel")
        self.root.geometry("1000x700")
        
        self.workbook: Optional[ExcelWorkbook] = None
        self.current_sheet: Optional[ExcelWorksheet] = None
        self.protocol_manager = ProtocolManager()
        
        self._create_widgets()
        self._create_menu()
    
    def _create_menu(self):
        """Создать меню приложения."""
        menubar = tk.Menu(self.root)
        self.root.config(menu=menubar)
        
        # Файл меню
        file_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Файл", menu=file_menu)
        file_menu.add_command(label="Открыть...", command=self.open_file, accelerator="Ctrl+O")
        file_menu.add_command(label="Создать новый", command=self.create_new_file, accelerator="Ctrl+N")
        file_menu.add_separator()
        file_menu.add_command(label="Сохранить", command=self.save_file, accelerator="Ctrl+S")
        file_menu.add_command(label="Сохранить как...", command=self.save_file_as)
        file_menu.add_separator()
        file_menu.add_command(label="Закрыть файл", command=self.close_file)
        file_menu.add_separator()
        file_menu.add_command(label="Выход", command=self.root.quit, accelerator="Alt+F4")
        
        # Протокол меню
        protocol_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Протокол", menu=protocol_menu)
        protocol_menu.add_command(label="Новый протокол", command=self.new_protocol)
        protocol_menu.add_command(label="Открыть протокол", command=self.open_protocol)
        protocol_menu.add_command(label="Редактировать протокол", command=self.edit_protocol)
        
        # Листы меню
        sheets_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Листы", menu=sheets_menu)
        sheets_menu.add_command(label="Создать лист...", command=self.create_sheet)
        sheets_menu.add_separator()
        self.sheets_submenu = tk.Menu(sheets_menu, tearoff=0)
        sheets_menu.add_cascade(label="Перейти на лист", menu=self.sheets_submenu)
        
        # Правка меню (без изменений)
        edit_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Правка", menu=edit_menu)
        edit_menu.add_command(label="Добавить строку", command=self.add_row)
        edit_menu.add_command(label="Удалить строку", command=self.delete_row)
        edit_menu.add_separator()
        edit_menu.add_command(label="Сортировать по столбцу...", command=self.sort_column)
        
        # Горячие клавиши
        self.root.bind('<Control-o>', lambda e: self.open_file())
        self.root.bind('<Control-n>', lambda e: self.create_new_file())
        self.root.bind('<Control-s>', lambda e: self.save_file())
    
    def _create_widgets(self):
        """Создать виджеты интерфейса."""
        # Верхняя панель с информацией
        info_frame = ttk.Frame(self.root)
        info_frame.pack(fill=tk.X, padx=5, pady=5)
        
        ttk.Label(info_frame, text="Файл:").pack(side=tk.LEFT)
        self.file_label = ttk.Label(info_frame, text="Файл не открыт", foreground="gray")
        self.file_label.pack(side=tk.LEFT, padx=5)
        
        ttk.Label(info_frame, text="|").pack(side=tk.LEFT)
        
        ttk.Label(info_frame, text="Лист:").pack(side=tk.LEFT)
        self.sheet_label = ttk.Label(info_frame, text="")
        self.sheet_label.pack(side=tk.LEFT, padx=5)
        
        ttk.Label(info_frame, text="|").pack(side=tk.LEFT)
        
        ttk.Label(info_frame, text="Протокол:").pack(side=tk.LEFT)
        self.protocol_label = ttk.Label(info_frame, text="Не загружен", foreground="gray")
        self.protocol_label.pack(side=tk.LEFT, padx=5)
        
        # Панель кнопок
        button_frame = ttk.Frame(self.root)
        button_frame.pack(fill=tk.X, padx=5, pady=5)
        
        ttk.Button(button_frame, text="Новый диапазон", command=self.on_new_range).pack(side=tk.LEFT, padx=2)
        ttk.Button(button_frame, text="Добавить линию", command=self.on_add_line).pack(side=tk.LEFT, padx=2)
        ttk.Button(button_frame, text="Внести данные", command=self.on_insert_data).pack(side=tk.LEFT, padx=2)
        
        # Таблица данных
        table_frame = ttk.Frame(self.root)
        table_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        columns = ("col1", "col2", "col3", "col4", "col5", "col6", "col7", "col8", "col9", "col10")
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=20)
        
        for i, col in enumerate(columns, 1):
            self.tree.heading(col, text=f"Столбец {i}")
            self.tree.column(col, width=100)
        
        scrollbar_y = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=self.tree.yview)
        scrollbar_x = ttk.Scrollbar(table_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=scrollbar_y.set, xscrollcommand=scrollbar_x.set)
        
        self.tree.grid(row=0, column=0, sticky="nsew")
        scrollbar_y.grid(row=0, column=1, sticky="ns")
        scrollbar_x.grid(row=1, column=0, sticky="ew")
        
        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)
        
        # Статус бар
        self.status_var = tk.StringVar(value="Готов")
        status_bar = ttk.Label(self.root, textvariable=self.status_var, relief=tk.SUNKEN, anchor=tk.W)
        status_bar.pack(fill=tk.X, side=tk.BOTTOM)
    
    def update_status(self, message: str):
        """Обновить статус бар."""
        self.status_var.set(message)
    
    def refresh_table(self):
        """Обновить таблицу данными из текущего листа."""
        for item in self.tree.get_children():
            self.tree.delete(item)
        
        if not self.current_sheet:
            return
        
        try:
            data = self.current_sheet.get_all_data()
            for row in data:
                values = [str(cell) if cell is not None else "" for cell in row[:10]]
                while len(values) < 10:
                    values.append("")
                self.tree.insert("", tk.END, values=values)
        except Exception as e:
            self.update_status(f"Ошибка обновления таблицы: {e}")
    
    def update_sheets_menu(self):
        """Обновить меню листов."""
        self.sheets_submenu.delete(0, tk.END)
        
        if not self.workbook:
            return
        
        for sheet_name in self.workbook.get_sheet_names():
            self.sheets_submenu.add_command(
                label=sheet_name,
                command=lambda name=sheet_name: self.switch_sheet(name)
            )
    
    # Методы работы с файлами
    def open_file(self):
        """Открыть существующий файл."""
        file_path = filedialog.askopenfilename(
            title="Открыть файл Excel",
            filetypes=[("Excel files", "*.xlsx")]
        )
        
        if not file_path:
            return
        
        try:
            if self.workbook:
                self.workbook.close()
            
            self.workbook = ExcelWorkbook(file_path)
            self.current_sheet = self.workbook.get_active_sheet()
            
            self.file_label.config(text=f"Файл: {file_path}", foreground="green")
            self.sheet_label.config(text=f"Лист: {self.current_sheet.name}")
            
            self.refresh_table()
            self.update_sheets_menu()
            self.update_status(f"Открыт файл: {file_path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось открыть файл:\n{str(e)}")
            self.update_status("Ошибка открытия файла")
    
    def create_new_file(self):
        """Создать новый файл."""
        if self.workbook:
            if not messagebox.askyesno("Подтверждение", "Текущий файл будет закрыт. Продолжить?"):
                return
            self.workbook.close()
        
        self.workbook = ExcelWorkbook()
        self.current_sheet = self.workbook.get_active_sheet()
        
        self.file_label.config(text="Новый файл (не сохранен)", foreground="orange")
        self.sheet_label.config(text=f"Лист: {self.current_sheet.name}")
        
        self.refresh_table()
        self.update_sheets_menu()
        self.update_status("Создан новый файл")
    
    def save_file(self):
        """Сохранить текущий файл."""
        if not self.workbook:
            messagebox.showwarning("Предупреждение", "Нет открытого файла")
            return
        
        if not self.workbook.file_path:
            self.save_file_as()
            return
        
        try:
            self.workbook.save()
            self.file_label.config(text=f"Файл: {self.workbook.file_path}", foreground="green")
            self.update_status("Файл сохранен")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить файл:\n{str(e)}")
            self.update_status("Ошибка сохранения")
    
    def save_file_as(self):
        """Сохранить файл как..."""
        if not self.workbook:
            return
        
        file_path = filedialog.asksaveasfilename(
            title="Сохранить файл Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")]
        )
        
        if not file_path:
            return
        
        try:
            self.workbook.save(file_path)
            self.file_label.config(text=f"Файл: {file_path}", foreground="green")
            self.update_status(f"Файл сохранен: {file_path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить файл:\n{str(e)}")
    
    def close_file(self):
        """Закрыть текущий файл."""
        if not self.workbook:
            return
        
        if self.workbook.file_path and self.workbook._lock:
            if not messagebox.askyesno("Подтверждение", "Закрыть текущий файл?"):
                return
        
        self.workbook.close()
        self.workbook = None
        self.current_sheet = None
        
        self.file_label.config(text="Файл не открыт", foreground="gray")
        self.sheet_label.config(text="")
        
        for item in self.tree.get_children():
            self.tree.delete(item)
        
        self.update_status("Файл закрыт")
    
    def create_sheet(self):
        """Создать новый лист."""
        if not self.workbook:
            messagebox.showwarning("Предупреждение", "Сначала откройте или создайте файл")
            return
        
        dialog = tk.Toplevel(self.root)
        dialog.title("Создать лист")
        dialog.geometry("300x120")
        dialog.transient(self.root)
        dialog.grab_set()
        
        ttk.Label(dialog, text="Имя листа:").pack(pady=10)
        name_entry = ttk.Entry(dialog, width=40)
        name_entry.pack(pady=5)
        name_entry.focus()
        
        def on_create():
            name = name_entry.get().strip()
            if not name:
                messagebox.showwarning("Предупреждение", "Введите имя листа", parent=dialog)
                return
            
            try:
                self.workbook.create_sheet(name)
                self.update_sheets_menu()
                self.update_status(f"Создан лист: {name}")
                dialog.destroy()
            except Exception as e:
                messagebox.showerror("Ошибка", f"Не удалось создать лист:\n{str(e)}", parent=dialog)
        
        ttk.Button(dialog, text="Создать", command=on_create).pack(pady=5)
    
    def switch_sheet(self, sheet_name: str):
        """Переключиться на другой лист."""
        if not self.workbook:
            return
        
        try:
            self.workbook.set_active_sheet(sheet_name)
            self.current_sheet = self.workbook.get_active_sheet()
            self.sheet_label.config(text=f"Лист: {self.current_sheet.name}")
            self.refresh_table()
            self.update_status(f"Переключен на лист: {sheet_name}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось переключить лист:\n{str(e)}")
    
    def add_row(self):
        """Добавить строку."""
        if not self.current_sheet:
            messagebox.showwarning("Предупреждение", "Нет активного листа")
            return
        
        dialog = tk.Toplevel(self.root)
        dialog.title("Добавить строку")
        dialog.geometry("400x200")
        dialog.transient(self.root)
        dialog.grab_set()
        
        ttk.Label(dialog, text="Введите значения через запятую:").pack(pady=10)
        values_entry = ttk.Entry(dialog, width=50)
        values_entry.pack(pady=5)
        values_entry.focus()
        
        def on_add():
            values_str = values_entry.get().strip()
            if not values_str:
                messagebox.showwarning("Предупреждение", "Введите значения", parent=dialog)
                return
            
            values = [v.strip() for v in values_str.split(',')]
            
            try:
                self.current_sheet.add_row(values)
                self.refresh_table()
                self.update_status("Строка добавлена")
                dialog.destroy()
            except RowValidationError as e:
                messagebox.showerror("Ошибка валидации", str(e), parent=dialog)
            except Exception as e:
                messagebox.showerror("Ошибка", f"Не удалось добавить строку:\n{str(e)}", parent=dialog)
        
        ttk.Button(dialog, text="Добавить", command=on_add).pack(pady=10)
    
    def delete_row(self):
        """Удалить выбранную строку."""
        selection = self.tree.selection()
        if not selection:
            messagebox.showwarning("Предупреждение", "Выберите строку для удаления")
            return
        
        if not self.current_sheet:
            return
        
        if not messagebox.askyesno("Подтверждение", "Удалить выбранную строку?"):
            return
        
        item = selection[0]
        row_index = self.tree.index(item) + 2  # +2 потому что индексация с 1 и есть заголовок
        
        try:
            self.current_sheet.delete_row(row_index)
            self.refresh_table()
            self.update_status("Строка удалена")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось удалить строку:\n{str(e)}")
    
    def sort_column(self):
        """Сортировать по столбцу."""
        if not self.current_sheet:
            messagebox.showwarning("Предупреждение", "Нет активного листа")
            return
        
        dialog = tk.Toplevel(self.root)
        dialog.title("Сортировка")
        dialog.geometry("250x150")
        dialog.transient(self.root)
        dialog.grab_set()
        
        ttk.Label(dialog, text="Номер столбца:").pack(pady=10)
        col_entry = ttk.Entry(dialog, width=10)
        col_entry.pack(pady=5)
        col_entry.focus()
        
        ttk.Label(dialog, text="(1-10)").pack()
        
        def on_sort():
            try:
                col_num = int(col_entry.get().strip())
                if col_num < 1 or col_num > 10:
                    raise ValueError("Номер столбца должен быть от 1 до 10")
                
                self.current_sheet.sort_by_column(col_num)
                self.refresh_table()
                self.update_status(f"Сортировка по столбцу {col_num}")
                dialog.destroy()
            except ValueError as e:
                messagebox.showerror("Ошибка", str(e), parent=dialog)
            except Exception as e:
                messagebox.showerror("Ошибка", f"Не удалось сортировать:\n{str(e)}", parent=dialog)
        
        ttk.Button(dialog, text="Сортировать", command=on_sort).pack(pady=10)
    
    # Методы работы с протоколом
    def new_protocol(self):
        """Создать новый протокол."""
        name = simpledialog.askstring(
            "Новый протокол",
            "Введите название протокола:",
            initialvalue=self.protocol_manager.get_new_protocol_name()
        )
        
        if name is None:
            return
        
        if not name.strip():
            name = self.protocol_manager.get_new_protocol_name()
        
        try:
            self.protocol_manager.create_protocol(name.strip())
            self.protocol_label.config(text=f"Протокол: {name}", foreground="green")
            self.update_status(f"Создан протокол: {name}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать протокол:\n{str(e)}")
    
    def open_protocol(self):
        """Открыть существующий протокол."""
        file_path = filedialog.askopenfilename(
            title="Открыть протокол",
            filetypes=[("JSON files", "*.json")]
        )
        
        if not file_path:
            return
        
        try:
            self.protocol_manager.open_protocol(file_path)
            protocol_name = os.path.basename(file_path).replace('.json', '')
            self.protocol_label.config(text=f"Протокол: {protocol_name}", foreground="green")
            self.update_status(f"Открыт протокол: {protocol_name}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось открыть протокол:\n{str(e)}")
    
    def edit_protocol(self):
        """Редактировать текущий протокол."""
        if not self.protocol_manager.current_protocol:
            messagebox.showwarning("Предупреждение", "Нет открытого протокола")
            return
        
        # Простое отображение информации о протоколе
        info = f"Протокол: {self.protocol_manager.current_protocol['name']}\n\n"
        info += f"Диапазоны: {len(self.protocol_manager.current_protocol['ranges'])}\n"
        info += f"Линии: {len(self.protocol_manager.current_protocol['lines'])}\n\n"
        
        if self.protocol_manager.current_protocol['lines']:
            info += "Линии:\n"
            for line in self.protocol_manager.current_protocol['lines']:
                info += f"  - {line['particle']}: w0={line['w0']}\n"
        
        messagebox.showinfo("Информация о протоколе", info)
    
    # Методы кнопок
    def on_new_range(self):
        """Обработчик кнопки 'Новый диапазон'."""
        if not self.protocol_manager.current_protocol:
            messagebox.showwarning("Предупреждение", "Сначала создайте или откройте протокол")
            return
        
        name = simpledialog.askstring("Новый диапазон", "Введите название диапазона:")
        if name:
            try:
                self.protocol_manager.add_range(name.strip())
                self.update_status(f"Добавлен диапазон: {name}")
            except Exception as e:
                messagebox.showerror("Ошибка", f"Не удалось добавить диапазон:\n{str(e)}")
    
    def on_add_line(self):
        """Обработчик кнопки 'Добавить линию'."""
        if not self.protocol_manager.current_protocol:
            messagebox.showwarning("Предупреждение", "Сначала создайте или откройте протокол")
            return
        
        ranges = self.protocol_manager.current_protocol.get('ranges', [])
        if not ranges:
            messagebox.showwarning("Предупреждение", "Сначала создайте диапазон")
            return
        
        # Диалог выбора диапазона
        dialog = tk.Toplevel(self.root)
        dialog.title("Добавить линию")
        dialog.geometry("350x200")
        dialog.transient(self.root)
        dialog.grab_set()
        
        ttk.Label(dialog, text="Диапазон:").pack(pady=5)
        range_var = tk.StringVar(value=ranges[0]['name'])
        range_combo = ttk.Combobox(dialog, textvariable=range_var, width=30)
        range_combo['values'] = [r['name'] for r in ranges]
        range_combo.pack(pady=5)
        
        ttk.Label(dialog, text="Излучающая частица:").pack(pady=5)
        particle_entry = ttk.Entry(dialog, width=30)
        particle_entry.pack(pady=5)
        particle_entry.focus()
        
        ttk.Label(dialog, text="Длина волны w0:").pack(pady=5)
        w0_entry = ttk.Entry(dialog, width=30)
        w0_entry.pack(pady=5)
        
        def on_add():
            particle = particle_entry.get().strip()
            if not particle:
                messagebox.showwarning("Предупреждение", "Введите частицу", parent=dialog)
                return
            
            try:
                w0 = float(w0_entry.get().strip())
            except ValueError:
                messagebox.showerror("Ошибка", "w0 должно быть числом", parent=dialog)
                return
            
            try:
                self.protocol_manager.add_line(range_var.get(), particle, w0)
                self.update_status(f"Добавлена линия: {particle}, w0={w0}")
                dialog.destroy()
            except Exception as e:
                messagebox.showerror("Ошибка", f"Не удалось добавить линию:\n{str(e)}", parent=dialog)
        
        ttk.Button(dialog, text="Добавить", command=on_add).pack(pady=10)
    
    def on_insert_data(self):
        """Обработчик кнопки 'Внести данные'."""
        if not self.workbook:
            messagebox.showwarning("Предупреждение", "Сначала откройте или создайте файл Excel")
            return
        
        if not self.protocol_manager.current_protocol:
            messagebox.showwarning("Предупреждение", "Сначала загрузите протокол")
            return
        
        # Диалог ввода данных
        dialog = tk.Toplevel(self.root)
        dialog.title("Внести данные")
        dialog.geometry("400x250")
        dialog.transient(self.root)
        dialog.grab_set()
        
        ttk.Label(dialog, text="Введите данные в формате:").pack(pady=5)
        ttk.Label(dialog, text="число<TAB>число", foreground="gray").pack()
        ttk.Label(dialog, text="число<TAB>число", foreground="gray").pack()
        ttk.Label(dialog, text="", text="Пример:").pack()
        ttk.Label(dialog, text="499,68625<TAB>0,01685", foreground="blue").pack()
        ttk.Label(dialog, text="2,32556<TAB>0,03365", foreground="blue").pack()
        
        text_widget = tk.Text(dialog, width=40, height=6)
        text_widget.pack(pady=10)
        text_widget.focus()
        
        def on_submit():
            content = text_widget.get("1.0", tk.END).strip()
            lines = content.split('\n')
            
            if len(lines) != 2:
                messagebox.showerror("Ошибка", "Должно быть ровно 2 строки", parent=dialog)
                return
            
            try:
                # Парсим первую строку: w, dw
                parts1 = lines[0].split('\t')
                if len(parts1) != 2:
                    raise ValueError("Первая строка должна содержать 2 числа через табуляцию")
                w = float(parts1[0].replace(',', '.'))
                dw = float(parts1[1].replace(',', '.'))
                
                # Парсим вторую строку: fwhm, dfwhm
                parts2 = lines[1].split('\t')
                if len(parts2) != 2:
                    raise ValueError("Вторая строка должна содержать 2 числа через табуляцию")
                fwhm = float(parts2[0].replace(',', '.'))
                dfwhm = float(parts2[1].replace(',', '.'))
                
            except ValueError as e:
                messagebox.showerror("Ошибка", f"Ошибка парсинга чисел: {e}", parent=dialog)
                return
            
            # Ищем ближайшую линию
            nearest_line = self.protocol_manager.find_nearest_line(w, max_distance=1.0)
            
            if not nearest_line:
                messagebox.showerror("Ошибка", f"Не найдена линия в диапазоне ±1.0 от w={w}", parent=dialog)
                dialog.destroy()
                return
            
            dialog.destroy()
            
            # Запрашиваем delay
            delay_str = simpledialog.askstring(
                "Ввод delay",
                f"Найдена линия: {nearest_line['particle']}, w0={nearest_line['w0']}\nВведите delay:"
            )
            
            if delay_str is None:
                return
            
            try:
                delay = float(delay_str.replace(',', '.'))
            except ValueError:
                messagebox.showerror("Ошибка", "delay должно быть числом")
                return
            
            # Записываем данные в Excel
            self._write_spectrum_data(nearest_line, delay, w, dw, fwhm, dfwhm)
        
        ttk.Button(dialog, text="Внести", command=on_submit).pack(pady=10)
    
    def _write_spectrum_data(self, line: Dict, delay: float, w: float, dw: float, fwhm: float, dfwhm: float):
        """Записать данные спектра в Excel."""
        particle = line['particle']
        w0 = line['w0']
        
        # Переключаемся на лист с названием частицы
        sheet_names = self.workbook.get_sheet_names()
        if particle not in sheet_names:
            self.workbook.create_sheet(particle)
        
        self.workbook.set_active_sheet(particle)
        sheet = self.workbook.get_active_sheet()
        
        # Проверяем заголовки
        headers = ["w0", "delay", "w", "dw", "fwhm", "dfwhm", "shift", "dshift"]
        existing_headers = sheet.get_row(1)
        
        if not existing_headers or existing_headers[0] != "w0":
            # Добавляем заголовки
            sheet.add_row(headers)
        
        # Находим номер последней строки
        all_data = sheet.get_all_data()
        row_num = len(all_data) + 1
        
        # Рассчитываем формулы для shift и dshift
        # shift = (w - w0) * 1000 = (C{row} - A{row}) * 1000
        # dshift = dw * 1000 = D{row} * 1000
        shift_formula = f"=(C{row_num}-A{row_num})*1000"
        dshift_formula = f"=D{row_num}*1000"
        
        # Данные для строки
        row_data = [w0, delay, w, dw, fwhm, dfwhm, shift_formula, dshift_formula]
        
        # Добавляем строку
        sheet.add_row(row_data)
        
        self.refresh_table()
        self.update_status(f"Данные внесены для {particle}: w0={w0}, delay={delay}")


def main():
    root = tk.Tk()
    app = SpectrumExcelApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
